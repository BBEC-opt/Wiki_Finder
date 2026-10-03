from pathlib import Path

import pytest

from app.parsing.service import ParsingService


@pytest.mark.asyncio
async def test_markdown_parse_writes_manifest(tmp_path: Path):
    source = tmp_path / "guide.md"
    source.write_text("# 指南\n\n正文\n\n|A|B|\n|---|---|\n|1|2|", encoding="utf-8")
    parser = ParsingService(tmp_path / "artifacts")
    result = await parser.parse(source, "k1")
    assert result.parser_engine == "builtin"
    assert "# 指南" in result.markdown
    assert (tmp_path / "artifacts" / "k1" / "manifest.json").exists()


@pytest.mark.asyncio
async def test_docx_order_and_table(tmp_path: Path):
    from docx import Document
    source = tmp_path / "sample.docx"
    document = Document()
    document.add_heading("产品说明", 1)
    document.add_paragraph("正文内容")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "名称", "值"
    table.cell(1, 0).text, table.cell(1, 1).text = "模式", "自动"
    document.save(source)
    result = await ParsingService(tmp_path / "artifacts").parse(source, "k2")
    assert "# 产品说明" in result.markdown
    assert "| 名称 | 值 |" in result.markdown
    assert result.tables


@pytest.mark.asyncio
async def test_xlsx_preserves_sheet_and_table(tmp_path: Path):
    import openpyxl
    source = tmp_path / "sample.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "库存"
    sheet.append(["产品", "数量"])
    sheet.append(["星云中控", 8])
    book.save(source)
    result = await ParsingService(tmp_path / "artifacts").parse(source, "k3")
    assert "# 工作表：库存" in result.markdown
    assert "星云中控" in result.markdown
    assert result.tables


@pytest.mark.asyncio
async def test_pdf_retains_page_marker(tmp_path: Path):
    import fitz
    source = tmp_path / "sample.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Reset button: hold for ten seconds.")
    document.save(source)
    document.close()
    result = await ParsingService(tmp_path / "artifacts").parse(source, "k4")
    assert "<!-- page:1 -->" in result.markdown
    assert result.pages[0].page_number == 1


@pytest.mark.asyncio
async def test_pdf_inserts_extracted_image_references(tmp_path: Path):
    import fitz
    source = tmp_path / "image.pdf"
    image_path = tmp_path / "pixel.png"
    pixmap = fitz.Pixmap(fitz.csRGB, (0, 0, 2, 2), 0)
    pixmap.clear_with(0xFF3366)
    pixmap.save(image_path)
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "A document page with enough searchable text for parsing.")
    page.insert_image(fitz.Rect(72, 100, 172, 200), filename=image_path)
    document.save(source)
    document.close()

    result = await ParsingService(tmp_path / "artifacts").parse(source, "with-image")

    assert len(result.images) == 1
    assert result.images[0].markdown_ref in result.markdown
