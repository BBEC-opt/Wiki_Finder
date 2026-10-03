import asyncio
import hashlib
import io
import re
import time
from email import policy
from email.parser import BytesParser
from pathlib import Path
from uuid import uuid4

from charset_normalizer import from_bytes

from app.domain import ParsedDocument, ParsedImage, ParsedPage, ParsedTable, ParseWarning
from app.core.errors import ParseError
from app.parsing.core import ParseRequest, markdown_table, normalize_markdown


def _save_image(data: bytes, suffix: str, request: ParseRequest, page: int | None, name: str) -> ParsedImage:
    digest = hashlib.sha256(data).hexdigest()
    image_id = digest[:24]
    image_dir = request.artifact_dir / "images"
    image_dir.mkdir(parents=True, exist_ok=True)
    suffix = suffix.lower().lstrip(".") or "bin"
    target = image_dir / f"{image_id}.{suffix}"
    if not target.exists():
        target.write_bytes(data)
    ref = f"artifact://{image_id}"
    return ParsedImage(
        id=image_id, original_name=name, mime_type=f"image/{suffix}", storage_path=str(target),
        markdown_ref=ref, page_number=page, content_hash=digest,
    )


class BuiltinParser:
    name = "builtin"
    supported_types = {
        "txt", "md", "markdown", "html", "htm", "mhtml", "pdf", "docx",
        "xlsx", "csv", "pptx", "png", "jpg", "jpeg", "webp", "bmp", "tiff",
    }

    def __init__(self, ocr_engine: str = "none"):
        self.ocr_engine = ocr_engine

    def available(self) -> tuple[bool, str]:
        return True, ""

    async def parse(self, request: ParseRequest, file_type: str) -> ParsedDocument:
        started = time.perf_counter()
        method = getattr(self, f"parse_{file_type}", None)
        if method is None and file_type in {"markdown"}:
            method = self.parse_md
        if method is None and file_type in {"htm"}:
            method = self.parse_html
        if method is None and file_type in {"jpeg", "webp", "bmp", "tiff"}:
            method = self.parse_image
        if method is None:
            raise ParseError("unsupported_file_type", f"builtin 不支持 {file_type}")
        result = await asyncio.to_thread(method, request)
        result.duration_ms = round((time.perf_counter() - started) * 1000)
        return result

    @staticmethod
    def _decode(path: Path) -> tuple[str, str]:
        raw = path.read_bytes()
        match = from_bytes(raw).best()
        if not match:
            raise ParseError("text_encoding_failed", "无法识别文本编码")
        return str(match), match.encoding or "unknown"

    def parse_txt(self, request: ParseRequest) -> ParsedDocument:
        text, encoding = self._decode(request.path)
        return ParsedDocument(normalize_markdown(text), self.name, {"encoding": encoding})

    def parse_md(self, request: ParseRequest) -> ParsedDocument:
        text, encoding = self._decode(request.path)
        return ParsedDocument(normalize_markdown(text), self.name, {"encoding": encoding, "format": "markdown"})

    def parse_html(self, request: ParseRequest) -> ParsedDocument:
        from bs4 import BeautifulSoup
        from markdownify import markdownify
        text, encoding = self._decode(request.path)
        soup = BeautifulSoup(text, "html.parser")
        for node in soup(["script", "style", "noscript", "nav"]):
            node.decompose()
        title = soup.title.get_text(strip=True) if soup.title else request.title
        body = soup.find("article") or soup.find("main") or soup.body or soup
        markdown = markdownify(str(body), heading_style="ATX")
        return ParsedDocument(normalize_markdown(markdown), self.name, {"encoding": encoding, "title": title})

    def parse_mhtml(self, request: ParseRequest) -> ParsedDocument:
        message = BytesParser(policy=policy.default).parsebytes(request.path.read_bytes())
        html = ""
        for part in message.walk():
            if part.get_content_type() == "text/html":
                html = part.get_content()
                break
        if not html:
            raise ParseError("empty_document", "MHTML 中没有 HTML 正文")
        from bs4 import BeautifulSoup
        from markdownify import markdownify
        soup = BeautifulSoup(html, "html.parser")
        for node in soup(["script", "style", "noscript"]):
            node.decompose()
        return ParsedDocument(normalize_markdown(markdownify(str(soup), heading_style="ATX")), self.name, {"format": "mhtml"})

    def parse_docx(self, request: ParseRequest) -> ParsedDocument:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        doc = Document(request.path)
        blocks, images, tables = [], [], []
        rel_images = {}
        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                blob = rel.target_part.blob
                suffix = rel.target_part.partname.rsplit(".", 1)[-1]
                image = _save_image(blob, suffix, request, None, Path(rel.target_part.partname).name)
                images.append(image)
                rel_images[rel.rId] = image
        for child in doc.element.body.iterchildren():
            if child.tag.endswith("}p"):
                paragraph = Paragraph(child, doc)
                text = paragraph.text.strip()
                style = paragraph.style.name if paragraph.style else ""
                if text:
                    match = re.match(r"Heading\s+(\d+)", style, re.I)
                    blocks.append(f"{'#' * min(6, int(match.group(1)))} {text}" if match else text)
                for drawing in child.xpath(".//a:blip"):
                    image = rel_images.get(drawing.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"))
                    if image:
                        blocks.append(f"![{image.original_name}]({image.markdown_ref})")
            elif child.tag.endswith("}tbl"):
                table = Table(child, doc)
                rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
                md = markdown_table(rows)
                if md:
                    item = ParsedTable(str(uuid4()), md, len(rows), max(map(len, rows), default=0))
                    tables.append(item)
                    blocks.append(md)
        return ParsedDocument(normalize_markdown("\n\n".join(blocks)), self.name,
                              {"paragraph_count": len(doc.paragraphs)}, images=images, tables=tables)

    def parse_xlsx(self, request: ParseRequest) -> ParsedDocument:
        import openpyxl
        book = openpyxl.load_workbook(request.path, data_only=True, read_only=False)
        blocks, tables = [], []
        for sheet in book.worksheets:
            blocks.append(f"# 工作表：{sheet.title}")
            rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
            while rows and not any(value not in (None, "") for value in rows[-1]):
                rows.pop()
            if not rows:
                continue
            width = max((i + 1 for row in rows for i, value in enumerate(row) if value not in (None, "")), default=0)
            rows = [row[:width] for row in rows]
            for offset in range(0, len(rows), 200):
                window = rows[offset:offset + 200]
                if offset and rows:
                    window = [rows[0], *window]
                md = markdown_table(window)
                tables.append(ParsedTable(str(uuid4()), md, len(window), width))
                blocks.append(f"## 行 {offset + 1}-{min(len(rows), offset + 200)}\n\n{md}")
        return ParsedDocument(normalize_markdown("\n\n".join(blocks)), self.name,
                              {"sheet_count": len(book.sheetnames)}, tables=tables)

    def parse_csv(self, request: ParseRequest) -> ParsedDocument:
        import csv
        text, encoding = self._decode(request.path)
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
        blocks, tables = [], []
        for offset in range(0, len(rows), 200):
            window = rows[offset:offset + 200]
            if offset and rows:
                window = [rows[0], *window]
            md = markdown_table(window)
            tables.append(ParsedTable(str(uuid4()), md, len(window), max(map(len, window), default=0)))
            blocks.append(md)
        return ParsedDocument(normalize_markdown("\n\n".join(blocks)), self.name,
                              {"encoding": encoding, "delimiter": dialect.delimiter, "row_count": len(rows)}, tables=tables)

    def parse_pptx(self, request: ParseRequest) -> ParsedDocument:
        from pptx import Presentation
        presentation = Presentation(request.path)
        blocks, images, tables, pages = [], [], [], []
        for page_no, slide in enumerate(presentation.slides, 1):
            page_blocks = [f"## 第 {page_no} 页"]
            for shape in sorted(slide.shapes, key=lambda x: (getattr(x, "top", 0), getattr(x, "left", 0))):
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    page_blocks.append(shape.text.strip())
                if getattr(shape, "has_table", False):
                    rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                    md = markdown_table(rows)
                    tables.append(ParsedTable(str(uuid4()), md, len(rows), max(map(len, rows), default=0), page_no))
                    page_blocks.append(md)
                if getattr(shape, "shape_type", None) == 13:
                    image = shape.image
                    saved = _save_image(image.blob, image.ext, request, page_no, image.filename or f"slide-{page_no}.{image.ext}")
                    images.append(saved)
                    page_blocks.append(f"![{saved.original_name}]({saved.markdown_ref})")
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    page_blocks.append(f"### 备注\n\n{notes}")
            page_md = "\n\n".join(page_blocks)
            pages.append(ParsedPage(page_no, page_md))
            blocks.append(page_md)
        return ParsedDocument(normalize_markdown("\n\n".join(blocks)), self.name,
                              {"slide_count": len(presentation.slides)}, pages, images, tables)

    def parse_pdf(self, request: ParseRequest) -> ParsedDocument:
        import fitz
        try:
            document = fitz.open(request.path)
        except Exception as exc:
            raise ParseError("pdf_open_failed", f"PDF 打开失败：{exc}") from exc
        if document.needs_pass:
            raise ParseError("encrypted_pdf", "PDF 已加密，无法解析")
        pages, images, warnings, blocks = [], [], [], []
        page_image_markdown = {}
        low_text_pages = []
        for page_no, page in enumerate(document, 1):
            page_blocks = page.get_text("blocks", sort=True)
            text = "\n\n".join(str(block[4]).strip() for block in page_blocks if str(block[4]).strip())
            if len(text.strip()) < 40:
                low_text_pages.append(page_no)
            image_blocks = []
            for image_no, info in enumerate(page.get_images(full=True), 1):
                try:
                    extracted = document.extract_image(info[0])
                    saved = _save_image(extracted["image"], extracted.get("ext", "png"), request, page_no, f"page-{page_no}-{image_no}")
                    images.append(saved)
                    image_blocks.append(f"![{saved.original_name}]({saved.markdown_ref})")
                except Exception as exc:
                    warnings.append(ParseWarning("image_extract_failed", str(exc), page_no))
            page_image_markdown[page_no] = "\n\n".join(image_blocks)
            page_md = f"<!-- page:{page_no} -->\n\n{text}\n\n{page_image_markdown[page_no]}".strip()
            pages.append(ParsedPage(page_no, page_md))
            blocks.append(page_md)
        if low_text_pages:
            ocr_results = self._ocr_pdf_pages(document, low_text_pages, request, warnings)
            for page_no, ocr_text in ocr_results.items():
                pages[page_no - 1].markdown = f"<!-- page:{page_no} -->\n\n{ocr_text}\n\n{page_image_markdown[page_no]}".strip()
                blocks[page_no - 1] = pages[page_no - 1].markdown
        metadata = dict(document.metadata or {})
        metadata.update({"page_count": len(document), "ocr_page_count": len(low_text_pages) if self.ocr_engine != "none" else 0})
        if low_text_pages and self.ocr_engine == "none":
            warnings.append(ParseWarning("scanned_pages_without_ocr", f"{len(low_text_pages)} 页疑似扫描件，未配置 OCR"))
        document.close()
        return ParsedDocument(normalize_markdown("\n\n".join(blocks)), self.name, metadata, pages, images, warnings=warnings)

    def _ocr_pdf_pages(self, document, page_numbers, request, warnings) -> dict[int, str]:
        if self.ocr_engine != "tesseract":
            return {}
        try:
            import pytesseract
            from PIL import Image
        except ImportError:
            warnings.append(ParseWarning("ocr_dependency_missing", "请安装 Pillow、pytesseract 和 Tesseract"))
            return {}
        results = {}
        for page_no in page_numbers:
            try:
                pix = document[page_no - 1].get_pixmap(dpi=200, alpha=False)
                image = Image.open(io.BytesIO(pix.tobytes("jpeg")))
                results[page_no] = pytesseract.image_to_string(image, lang="chi_sim+eng").strip()
            except Exception as exc:
                warnings.append(ParseWarning("ocr_failed", str(exc), page_no))
        return results

    def parse_png(self, request):
        return self.parse_image(request)

    def parse_jpg(self, request):
        return self.parse_image(request)

    def parse_image(self, request: ParseRequest) -> ParsedDocument:
        data = request.path.read_bytes()
        saved = _save_image(data, request.path.suffix, request, 1, request.path.name)
        warnings = []
        text = ""
        if self.ocr_engine == "tesseract":
            try:
                import pytesseract
                from PIL import Image
                text = pytesseract.image_to_string(Image.open(request.path), lang="chi_sim+eng").strip()
            except Exception as exc:
                warnings.append(ParseWarning("ocr_failed", str(exc), 1))
        else:
            warnings.append(ParseWarning("image_without_ocr", "图片已保存，但未配置 OCR", 1))
        markdown = f"![{request.path.name}]({saved.markdown_ref})\n\n{text}".strip()
        return ParsedDocument(markdown, self.name, {"page_count": 1}, [ParsedPage(1, markdown)], [saved], warnings=warnings)
